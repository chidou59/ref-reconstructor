"""
文件路径: core/doc_renderer.py
=========================================================
【功能】
负责生成最终的 Word 文档。
【核心升级】
- 支持自定义排版格式（字体、字号、行间距）
- 修复了中文 "宋体" 等字体在 Word 中不生效的问题 (需设置 eastAsia)
- 【关键修复 V6.11】修复中文文档缺失 "Heading 1" 样式导致的 KeyError 崩溃问题。
- 【关键修复 V6.12】智能保留用户原有的“参考文献”标题格式，仅替换内容。
=========================================================
"""

import random
import string
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from .utils import CITATION_PATTERN, parse_citation_ranges, BIB_SECTION_KEYWORDS, is_likely_section_header


class DocRenderer:
    def __init__(self, input_path):
        self.doc = Document(input_path)
        self._bookmark_id_counter = 10000000
        self._session_suffix = ''.join(random.choices(string.ascii_lowercase + string.digits, k=4))

    def render(self, output_path, id_mapping, new_bib_list, use_superscript=False, style_config=None):
        """
        执行渲染流程
        """
        # 1. 替换正文中的引用
        self._replace_all_citations(id_mapping, use_superscript)

        # 2. 原位替换参考文献章节 (Smart Replace)
        self._smart_replace_bibliography(new_bib_list, style_config)

        # 3. 保存
        self.doc.save(output_path)

    def _replace_all_citations(self, mapping, use_superscript=False):
        """遍历文档所有内容执行替换"""
        for para in self.doc.paragraphs:
            self._flatten_fields(para)
            self._process_paragraph_safe(para, mapping, use_superscript)

        for table in self.doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        self._flatten_fields(para)
                        self._process_paragraph_safe(para, mapping, use_superscript)

    def _smart_replace_bibliography(self, bib_list, style_config=None):
        """
        智能替换参考文献，并应用排版样式
        【逻辑更新】如果存在旧标题，则保留旧标题，仅替换下方内容。
        """
        paras = self.doc.paragraphs
        start_idx = -1
        end_idx = len(paras)

        # 1. 寻找开始位置 (标题行)
        for i in range(len(paras) - 1, -1, -1):
            text = paras[i].text.strip().lower()
            if len(text) < 20:
                for kw in BIB_SECTION_KEYWORDS:
                    if kw in text:
                        start_idx = i
                        break
            if start_idx != -1:
                break

        # 情况 A: 没找到旧标题 -> 在文末追加 (需要新建标题)
        if start_idx == -1:
            self._append_bibliography_at_end(bib_list, style_config, insert_header=True)
            return

        # 2. 寻找结束位置
        for i in range(start_idx + 1, len(paras)):
            text = paras[i].text.strip()
            if is_likely_section_header(text):
                end_idx = i
                break

        # 情况 B: 找到了旧标题 -> 保留标题(start_idx)，只删除它后面的内容
        # print(f"定位参考文献区: 标题在行 {start_idx}, 内容从 {start_idx+1} -> {end_idx}")

        # 3. 删除旧段落 (从标题的下一行开始删)
        paras_to_delete = []
        # 注意：这里从 start_idx + 1 开始，意味着保留 start_idx (即原标题)
        for i in range(start_idx + 1, end_idx):
            paras_to_delete.append(paras[i])

        for p in paras_to_delete:
            p._element.getparent().remove(p._element)

        # 4. 插入新内容
        # 删除后，原本在 end_idx 的段落现在跑到了 start_idx + 1 的位置
        current_paras = self.doc.paragraphs

        if start_idx + 1 < len(current_paras):
            # 如果后面还有内容（比如“附录”），就在它前面插
            insert_anchor = current_paras[start_idx + 1]
            # 因为保留了原标题，所以 insert_header=False
            self._insert_bibliography_before(bib_list, insert_anchor, style_config, insert_header=False)
        else:
            # 如果后面没内容了，直接追加
            self._append_bibliography_at_end(bib_list, style_config, insert_header=False)

    def _apply_heading_style(self, paragraph):
        """
        【修复核心】安全地应用标题样式。
        """
        styles = self.doc.styles
        if 'Heading 1' in styles:
            paragraph.style = 'Heading 1'
            return
        if '标题 1' in styles:
            paragraph.style = '标题 1'
            return
        if 'heading 1' in styles:
            paragraph.style = 'heading 1'
            return
        try:
            paragraph.style = 'Normal'
        except:
            pass
        if not paragraph.runs:
            paragraph.add_run(paragraph.text)
            paragraph.text = ""
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(16)
            r = run._element
            rPr = r.get_or_add_rPr()
            rFonts = rPr.get_or_add_rFonts()
            rFonts.set(qn('w:eastAsia'), '黑体')

    def _insert_bibliography_before(self, bib_list, anchor_para, style_config=None, insert_header=True):
        """
        在指定段落之前插入参考文献
        :param insert_header: 是否插入“参考文献”这个大标题
        """
        # 1. 插入标题 (仅当需要时)
        if insert_header:
            header_p = anchor_para.insert_paragraph_before("参考文献")
            self._apply_heading_style(header_p)

        # 2. 插入文献条目
        for i, content in enumerate(bib_list):
            ref_id = i + 1
            p = anchor_para.insert_paragraph_before()
            self._fill_bib_paragraph(p, ref_id, content, style_config)

    def _append_bibliography_at_end(self, bib_list, style_config=None, insert_header=True):
        """
        在文末追加参考文献
        :param insert_header: 是否插入“参考文献”这个大标题
        """
        if insert_header:
            p = self.doc.add_paragraph('参考文献')
            self._apply_heading_style(p)

        for i, content in enumerate(bib_list):
            ref_id = i + 1
            p = self.doc.add_paragraph()
            self._fill_bib_paragraph(p, ref_id, content, style_config)

    def _fill_bib_paragraph(self, p, ref_id, content, style_config=None):
        """
        填充参考文献段落内容（含XML格式设置和自定义排版）
        """
        # === 0. 应用自定义排版样式 ===
        if style_config:
            if 'line_spacing' in style_config:
                p.paragraph_format.line_spacing = style_config['line_spacing']
            font_name = style_config.get('font_name', '宋体')
            font_size = style_config.get('font_size', 10.5)

            def apply_font(run):
                run.font.size = Pt(font_size)
                run.font.name = font_name
                r = run._element
                rPr = r.get_or_add_rPr()
                rFonts = rPr.get_or_add_rFonts()
                rFonts.set(qn('w:eastAsia'), font_name)
        else:
            def apply_font(run):
                pass

        unique_bm_name = f"Ref_Bib_{ref_id}_{self._session_suffix}"

        # === 1. 强制注入 XML 样式 (悬挂缩进) ===
        pPr = p._element.get_or_add_pPr()
        for tag in ['w:ind', 'w:tabs']:
            old = pPr.find(qn(tag))
            if old is not None: pPr.remove(old)

        new_ind = OxmlElement('w:ind')
        new_ind.set(qn('w:leftChars'), '0')
        new_ind.set(qn('w:hangingChars'), '250')
        new_ind.set(qn('w:left'), '0')
        new_ind.set(qn('w:hanging'), '525')
        pPr.append(new_ind)

        tabs = OxmlElement('w:tabs')
        tab = OxmlElement('w:tab')
        tab.set(qn('w:val'), 'left')
        tab.set(qn('w:pos'), '0')
        tabs.append(tab)
        pPr.append(tabs)

        # === 2. 内容构建 ===
        r1 = p.add_run("[")
        apply_font(r1)

        bm_start = OxmlElement('w:bookmarkStart')
        bm_start.set(qn('w:id'), str(self._bookmark_id_counter))
        bm_start.set(qn('w:name'), unique_bm_name)
        p._element.append(bm_start)

        self._append_seq_field_elements(p._element, ref_id)

        bm_end = OxmlElement('w:bookmarkEnd')
        bm_end.set(qn('w:id'), str(self._bookmark_id_counter))
        p._element.append(bm_end)
        self._bookmark_id_counter += 1

        r2 = p.add_run("]")
        apply_font(r2)

        r3 = p.add_run("\t")

        r4 = p.add_run(f"{content}")
        apply_font(r4)

    def _flatten_fields(self, para):
        """域清洗"""
        try:
            runs = para.runs
            for run in runs:
                r_element = run._element
                for tag in ['w:fldChar', 'w:instrText']:
                    for node in r_element.findall(qn(tag)):
                        r_element.remove(node)
        except Exception:
            pass

    def _process_paragraph_safe(self, para, mapping, use_superscript=False):
        full_text = para.text
        if not full_text.strip():
            return

        matches = list(CITATION_PATTERN.finditer(full_text))
        if not matches:
            return

        run_map = []
        current_pos = 0
        for run in para.runs:
            length = len(run.text)
            run_map.append((current_pos, current_pos + length, run))
            current_pos += length

        for m in reversed(matches):
            m_start, m_end = m.span()
            inner_str = m.group(1)

            old_ids = parse_citation_ranges(inner_str)
            new_ids = []
            for oid in old_ids:
                if oid in mapping:
                    new_ids.append(mapping[oid])

            if not new_ids:
                continue

            start_run_idx = -1
            end_run_idx = -1

            for i, (r_start, r_end, run) in enumerate(run_map):
                if r_start <= m_start < r_end:
                    start_run_idx = i
                if r_start < m_end <= r_end:
                    end_run_idx = i

            if start_run_idx == -1 or end_run_idx == -1:
                continue

            self._apply_replacement_with_ref_fields(para, run_map, start_run_idx, end_run_idx, m_start, m_end, new_ids,
                                                    use_superscript)

    def _apply_replacement_with_ref_fields(self, para, run_map, start_run_idx, end_run_idx, match_start, match_end,
                                           new_ids, use_superscript=False):
        start_r_start, _, start_run = run_map[start_run_idx]
        end_r_start, _, end_run = run_map[end_run_idx]

        cut_start = match_start - start_r_start
        cut_end = match_end - end_r_start

        prefix_text = start_run.text[:cut_start]
        suffix_text = end_run.text[cut_end:]

        start_run.text = prefix_text

        new_elements = self._build_citation_ref_fields(new_ids, start_run, use_superscript)

        parent = start_run._element.getparent()
        try:
            index = parent.index(start_run._element) + 1
        except ValueError:
            index = len(parent)

        for elem in reversed(new_elements):
            parent.insert(index, elem)

        last_inserted_element = new_elements[-1]

        if suffix_text:
            suffix_run = para.add_run(suffix_text)
            self._copy_run_format(end_run, suffix_run)
            parent.insert(parent.index(last_inserted_element) + 1, suffix_run._element)

        if start_run_idx != end_run_idx:
            end_run.text = ""
            for i in range(start_run_idx + 1, end_run_idx):
                run_map[i][2].text = ""

    def _build_citation_ref_fields(self, id_list, template_run, use_superscript=False):
        sorted_ids = sorted(list(set(id_list)))
        ranges = []
        if not sorted_ids: return []

        range_start = sorted_ids[0]
        prev = sorted_ids[0]
        for curr in sorted_ids[1:]:
            if curr == prev + 1:
                prev = curr
            else:
                ranges.append((range_start, prev))
                range_start = curr
                prev = curr
        ranges.append((range_start, prev))

        elements = []

        def apply_style_to_element(run_element):
            is_template_superscript = False
            if template_run.font and template_run.font.superscript:
                is_template_superscript = True

            if use_superscript or is_template_superscript:
                rPr = OxmlElement('w:rPr')
                vertAlign = OxmlElement('w:vertAlign')
                vertAlign.set(qn('w:val'), 'superscript')
                rPr.append(vertAlign)
                run_element.insert(0, rPr)

        def make_run(text):
            r = OxmlElement('w:r')
            apply_style_to_element(r)
            t = OxmlElement('w:t')
            t.text = text
            if text.strip() == "":
                t.set(qn('xml:space'), 'preserve')
            r.append(t)
            return r

        def make_ref_field(ref_id):
            unique_bm_name = f"Ref_Bib_{ref_id}_{self._session_suffix}"

            r_begin = OxmlElement('w:r')
            apply_style_to_element(r_begin)
            fldChar_begin = OxmlElement('w:fldChar')
            fldChar_begin.set(qn('w:fldCharType'), 'begin')
            r_begin.append(fldChar_begin)

            r_instr = OxmlElement('w:r')
            apply_style_to_element(r_instr)
            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = f' REF {unique_bm_name} \\h \\* CHARFORMAT '
            r_instr.append(instrText)

            r_sep = OxmlElement('w:r')
            apply_style_to_element(r_sep)
            fldChar_sep = OxmlElement('w:fldChar')
            fldChar_sep.set(qn('w:fldCharType'), 'separate')
            r_sep.append(fldChar_sep)

            r_res = make_run(str(ref_id))

            r_end = OxmlElement('w:r')
            apply_style_to_element(r_end)
            fldChar_end = OxmlElement('w:fldChar')
            fldChar_end.set(qn('w:fldCharType'), 'end')
            r_end.append(fldChar_end)

            return [r_begin, r_instr, r_sep, r_res, r_end]

        elements.append(make_run("["))
        for i, (start, end) in enumerate(ranges):
            if i > 0: elements.append(make_run(", "))

            if start == end:
                elements.extend(make_ref_field(start))
            elif start + 1 == end:
                elements.extend(make_ref_field(start))
                elements.append(make_run(", "))
                elements.extend(make_ref_field(end))
            else:
                elements.extend(make_ref_field(start))
                elements.append(make_run("-"))
                elements.extend(make_ref_field(end))

        elements.append(make_run("]"))
        return elements

    def _append_seq_field_elements(self, parent_element, ref_id):
        r_begin = OxmlElement('w:r')
        fldChar_begin = OxmlElement('w:fldChar')
        fldChar_begin.set(qn('w:fldCharType'), 'begin')
        r_begin.append(fldChar_begin)
        parent_element.append(r_begin)

        r_instr = OxmlElement('w:r')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = ' SEQ RefList '
        r_instr.append(instrText)
        parent_element.append(r_instr)

        r_sep = OxmlElement('w:r')
        fldChar_sep = OxmlElement('w:fldChar')
        fldChar_sep.set(qn('w:fldCharType'), 'separate')
        r_sep.append(fldChar_sep)
        parent_element.append(r_sep)

        r_res = OxmlElement('w:r')
        t_res = OxmlElement('w:t')
        t_res.text = str(ref_id)
        r_res.append(t_res)
        parent_element.append(r_res)

        r_end = OxmlElement('w:r')
        fldChar_end = OxmlElement('w:fldChar')
        fldChar_end.set(qn('w:fldCharType'), 'end')
        r_end.append(fldChar_end)
        parent_element.append(r_end)

    def _copy_run_format(self, source, target):
        try:
            f_src = source.font
            f_tgt = target.font
            if f_src.bold is not None: f_tgt.bold = f_src.bold
            if f_src.italic is not None: f_tgt.italic = f_src.italic
            if f_src.superscript is not None: f_tgt.superscript = f_src.superscript
        except Exception:
            pass